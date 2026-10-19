"""Colourful, image-rich PPTX built with python-pptx drawing primitives."""
import zlib
from typing import Optional
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from pydantic import BaseModel, Field
from PIL import Image

from ..config import OUTPUTS
from ..context import gather
from ..images import fetch_many
from ..llm import get_llm
from ..store import norm, save


class Slide(BaseModel):
    title: str = Field(description="short, catchy slide title")
    bullets: list[str] = Field(description="3-4 bullets, each under 14 words")
    notes: str = Field(description="speaker notes, 2-3 sentences")
    image_prompt: str = Field(default="", description="visual idea for an illustration, 6-12 words, no text in image")


class Deck(BaseModel):
    title: str
    subtitle: str = Field(description="one-line tagline")
    slides: list[Slide]


PALETTES = [  # dark, accent1, accent2, accent3, light background
    ("1E1B4B", "6366F1", "F97316", "10B981", "F5F5FF"),
    ("0F3D3E", "14B8A6", "F59E0B", "EC4899", "F0FDFA"),
    ("3B0A45", "D946EF", "FACC15", "38BDF8", "FDF4FF"),
    ("0B2545", "3B82F6", "F43F5E", "22C55E", "EFF6FF"),
]
W, H = 13.333, 7.5
FONT = "Calibri"


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def box(slide, shape, x, y, w, h, fill=None, line=None, radius=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(1.25)
    else:
        s.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    return s


def text(slide, x, y, w, h, t, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = t
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
    r.font.color.rgb = rgb(color)
    return tb


def picture(slide, path, x, y, w, h):
    """Add image cropped to fill the box (cover)."""
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    box_ratio, img_ratio = w / h, iw / ih
    if img_ratio > box_ratio:
        c = (1 - box_ratio / img_ratio) / 2
        pic.crop_left = pic.crop_right = c
    else:
        c = (1 - img_ratio / box_ratio) / 2
        pic.crop_top = pic.crop_bottom = c
    return pic


def framed_picture(slide, path, x, y, w, h, accent):
    box(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.18, y + 0.18, w, h, fill=accent, radius=0.04)
    pic = picture(slide, path, x, y, w, h)
    pic.line.color.rgb = rgb("FFFFFF")
    pic.line.width = Pt(4)


def background(slide, color):
    box(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=color)


def title_slide(prs, deck, pal, img, class_name):
    dark, a1, a2, a3, _ = pal
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, dark)
    box(s, MSO_SHAPE.OVAL, -1.5, -1.8, 5, 5, fill=a1)
    box(s, MSO_SHAPE.OVAL, 10.2, 5.3, 4.5, 4.5, fill=a2)
    box(s, MSO_SHAPE.OVAL, 6.4, -0.8, 1.4, 1.4, fill=a3)
    has_img = img is not None
    tw = 6.6 if has_img else 11
    text(s, 0.8, 1.0, 4, 0.5, class_name.upper(), 16, a2, bold=True)
    text(s, 0.8, 1.7, tw, 2.6, deck.title, 46, "FFFFFF", bold=True, anchor=MSO_ANCHOR.TOP)
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, 4.45, 1.6, 0.09, fill=a2, radius=0.5)
    text(s, 0.8, 4.7, tw, 1.2, deck.subtitle, 22, "E0E7FF", anchor=MSO_ANCHOR.TOP)
    if has_img:
        framed_picture(s, img, 7.9, 1.4, 4.9, 4.5, a2)
    return s


def content_slide(prs, sl, n, total, pal, img, class_name):
    dark, a1, a2, a3, light = pal
    accents = [a1, a2, a3]
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, light)
    box(s, MSO_SHAPE.RECTANGLE, 0, 0, W, 1.35, fill=dark)
    box(s, MSO_SHAPE.RECTANGLE, 0, 1.35, W, 0.07, fill=a2)
    box(s, MSO_SHAPE.OVAL, 0.45, 0.32, 0.72, 0.72, fill=a2)
    text(s, 0.45, 0.32, 0.72, 0.72, str(n), 24, "FFFFFF", bold=True, align=PP_ALIGN.CENTER)
    text(s, 1.4, 0.2, 11.3, 0.95, sl.title, 32, "FFFFFF", bold=True)

    img_left = n % 2 == 0
    body_w = 7.3 if img else 12.0
    body_x = (5.2 if img_left else 0.7) if img else 0.7
    if img:
        img_x = 0.7 if img_left else 8.2
        framed_picture(s, img, img_x, 1.95, 4.2, 4.5, a2 if img_left else a1)
    bullets = sl.bullets[:5]
    gap = 0.18
    ch = min(1.15, (5.1 - gap * (len(bullets) - 1)) / max(len(bullets), 1))
    y = 1.85
    for i, b in enumerate(bullets):
        c = accents[i % 3]
        box(s, MSO_SHAPE.ROUNDED_RECTANGLE, body_x, y, body_w, ch, fill="FFFFFF", line="E5E7EB", radius=0.2)
        box(s, MSO_SHAPE.ROUNDED_RECTANGLE, body_x, y, 0.14, ch, fill=c, radius=0.5)
        box(s, MSO_SHAPE.OVAL, body_x + 0.35, y + ch / 2 - 0.25, 0.5, 0.5, fill=c)
        text(s, body_x + 0.35, y + ch / 2 - 0.25, 0.5, 0.5, str(i + 1), 16, "FFFFFF", bold=True, align=PP_ALIGN.CENTER)
        text(s, body_x + 1.05, y, body_w - 1.2, ch, b, 19 if len(b) < 80 else 17, "1F2937")
        y += ch + gap
    text(s, 0.5, 6.95, 6, 0.4, class_name, 11, "9CA3AF")
    text(s, W - 1.5, 6.95, 1, 0.4, f"{n + 1} / {total}", 11, "9CA3AF", align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = sl.notes
    return s


def summary_slide(prs, sl, total, pal, class_name):
    dark, a1, a2, a3, _ = pal
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, dark)
    box(s, MSO_SHAPE.OVAL, 10.6, -1.5, 4.2, 4.2, fill=a1)
    box(s, MSO_SHAPE.OVAL, -1.2, 5.6, 3.6, 3.6, fill=a2)
    text(s, 0.8, 0.5, 11, 1.0, sl.title, 38, "FFFFFF", bold=True)
    box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, 1.5, 1.6, 0.09, fill=a2, radius=0.5)
    bullets = sl.bullets[:5]
    gap, y = 0.18, 1.95
    ch = min(0.95, (4.6 - gap * (len(bullets) - 1)) / max(len(bullets), 1))
    for i, b in enumerate(bullets):
        c = [a1, a2, a3][i % 3]
        box(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, y, 10.2, ch, fill="FFFFFF", radius=0.25)
        box(s, MSO_SHAPE.OVAL, 1.05, y + ch / 2 - 0.22, 0.44, 0.44, fill=c)
        text(s, 1.05, y + ch / 2 - 0.22, 0.44, 0.44, "✓", 16, "FFFFFF", bold=True, align=PP_ALIGN.CENTER)
        text(s, 1.7, y, 9.1, ch, b, 19, "1F2937")
        y += ch + gap
    text(s, W - 1.5, 6.95, 1, 0.4, f"{total} / {total}", 11, "C7D2FE", align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = sl.notes
    return s


def build_ppt(class_name: str, topics: list[str], num_slides: int, instructions: str) -> dict:
    context, sources = gather(topics)
    deck = get_llm().with_structured_output(Deck).invoke(
        f"Create an engaging classroom presentation for {class_name}.\n"
        f"Topics: {', '.join(topics)}\nExactly {num_slides} slides: slide 1 = title slide (title, subtitle; its own "
        f"bullets can be empty), middle slides = content, last slide = 'Key Takeaways' summary.\n"
        "Bullets must be short, clear and student-friendly. Give every slide an image_prompt describing a simple, "
        "colourful illustration of the idea.\n"
        f"Teacher instructions: {instructions or 'none'}\n\n"
        f"Use this source material (prefer it over your own knowledge):\n{context or '(none found)'}"
    )
    slides = deck.slides[:num_slides]
    imgs = fetch_many([s.image_prompt or f"{deck.title} {s.title}" for s in slides])
    pal = PALETTES[zlib.crc32(deck.title.encode()) % len(PALETTES)]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    total = len(slides)
    for i, (sl, img) in enumerate(zip(slides, imgs)):
        if i == 0:
            s = title_slide(prs, deck, pal, img, class_name)
            s.notes_slide.notes_text_frame.text = sl.notes
        elif i == total - 1 and total > 2:
            summary_slide(prs, sl, total, pal, class_name)
        else:
            content_slide(prs, sl, i, total, pal, img, class_name)
    fname = f"{norm(class_name)}-presentation.pptx"
    prs.save(OUTPUTS / fname)
    save(class_name, "ppt", {"file": fname, "topics": topics, "title": deck.title,
                             "slides": [s.model_dump() for s in slides]})
    return {"file": fname, "slides": total, "sources": sources, "images": sum(i is not None for i in imgs)}
